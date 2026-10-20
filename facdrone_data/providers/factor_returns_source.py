# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Marius Ciepluch
"""Provider transports for daily style-factor returns.

Two ``MarketDataClient`` adapters feed the same long ``factor_returns`` table. The
refresher passes the requested factor id through ``req.provider_symbol``:

* :class:`KenFrenchFactorReturnsClient` reads MKT/SMB/HML from the Dartmouth daily
  3-factor ZIP and MOM from the daily momentum ZIP. Ken French publishes percent
  returns, so the adapter converts them to decimals. Every Ken French daily file
  has the same shape — a preamble, one header row naming the columns, then
  eight-digit date rows — so one factor id is fully addressed by the archive it
  lives in and the raw column it occupies (``_KENFRENCH_SERIES``); the client
  holds one transport per archive.
* :class:`AqrFactorReturnsClient` reads QMJ/BAB from AQR daily workbooks. AQR
  already publishes decimal returns.

Both adapters return ``factor_id``/``date``/``ret`` with decimal daily simple
returns. Transport or parse faults return an empty frame so the refresher skips
that factor without taking down the whole download.
"""

from __future__ import annotations

import logging
from collections.abc import Callable, Iterator, Mapping
from dataclasses import dataclass
from typing import TYPE_CHECKING

import pandas as pd

from facdrone_data.providers.kenfrench_source import KenFrenchZip, _unzip_first_member
from facdrone_data.providers.policy import fetch_url
from facdrone_data_types.symbols import Provider

if TYPE_CHECKING:
    from facdrone_data_types.ports import FetchRequest

log = logging.getLogger(__name__)

_DATASET = "factor_returns"

_KENFRENCH_FTP = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp"
# Archive keys. One transport per archive; a factor names the archive it is in.
US_FF3 = "us_ff3"
US_MOMENTUM = "us_momentum"
EUROPE_FF3 = "europe_ff3"
EUROPE_MOMENTUM = "europe_momentum"
_ARCHIVE_URLS = {
    US_FF3: f"{_KENFRENCH_FTP}/F-F_Research_Data_Factors_daily_CSV.zip",
    US_MOMENTUM: f"{_KENFRENCH_FTP}/F-F_Momentum_Factor_daily_CSV.zip",
    EUROPE_FF3: f"{_KENFRENCH_FTP}/Europe_3_Factors_Daily_CSV.zip",
    EUROPE_MOMENTUM: f"{_KENFRENCH_FTP}/Europe_Mom_Factor_Daily_CSV.zip",
}


@dataclass(frozen=True)
class _KenFrenchSeries:
    """Where one factor id's raw percent column lives."""

    archive: str
    column: str


# The factor id to (archive, raw column) map. RF is a rate, not a factor return,
# so it is deliberately absent; MKT is the Mkt-RF excess column.
#
# The ``_EU`` ids are Ken French's European portfolios. They are the right
# EXPLANATORY MARKET for a book trading European instruments, but note what the
# suffix does not mean: Ken French states every regional series in US DOLLARS,
# so these are not EUR-denominated. That matters little for the self-financing
# spreads (SMB/HML/MOM are differences of two returns in one currency) and a
# great deal for MKT_EU, whose USD statement carries the full EURUSD move. The
# reporting registry decides what to do about that; ingestion just records the
# published series faithfully. The regional momentum column is ``WML``
# (winners-minus-losers), not the US files' ``Mom``.
_KENFRENCH_SERIES: dict[str, _KenFrenchSeries] = {
    "MKT": _KenFrenchSeries(US_FF3, "Mkt-RF"),
    "SMB": _KenFrenchSeries(US_FF3, "SMB"),
    "HML": _KenFrenchSeries(US_FF3, "HML"),
    "MOM": _KenFrenchSeries(US_MOMENTUM, "Mom"),
    "MKT_EU": _KenFrenchSeries(EUROPE_FF3, "Mkt-RF"),
    "SMB_EU": _KenFrenchSeries(EUROPE_FF3, "SMB"),
    "HML_EU": _KenFrenchSeries(EUROPE_FF3, "HML"),
    "MOM_EU": _KenFrenchSeries(EUROPE_MOMENTUM, "WML"),
}

# Ken French missing-data sentinels (percent units). Checked BEFORE the percent→decimal
# divide so a missing day becomes NaN (dropped by the normalizer), never a real
# ~-99.99% / -999% daily return.
_KENFRENCH_SENTINELS = (-99.99, -999.0)


def _percent_to_decimal(values: pd.Series) -> pd.Series:
    """Percent→decimal with the Ken French sentinels nulled first:
    ``-99.99``/``-999`` → NaN before the divide, else ``v/100``."""
    numeric = pd.to_numeric(values, errors="coerce")
    return (numeric / 100.0).where(~numeric.isin(_KENFRENCH_SENTINELS), other=None)


# AQR daily workbooks by factor id. QMJ and BAB are separate files; use the USA column.
_AQR_URLS = {
    "QMJ": "https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/"
    "Quality-Minus-Junk-Factors-Daily.xlsx",
    "BAB": "https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/"
    "Betting-Against-Beta-Equity-Factors-Daily.xlsx",
}
# AQR's CDN rejects the default urllib agent.
_AQR_USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120 Safari/537.36"
)

# Fetch one AQR workbook as bytes. Tests inject this boundary.
AqrFactorFile = Callable[[str], bytes]


def _window(frame: pd.DataFrame, req: FetchRequest) -> pd.DataFrame:
    """Return rows inside the inclusive request window."""
    if frame.empty:
        return frame
    dates = pd.to_datetime(frame["date"]).dt.date
    keep = (dates >= req.start) & (dates <= req.end)
    return frame[keep].reset_index(drop=True)


def _parse_column_daily(raw_csv: str, column: str) -> pd.DataFrame:
    """Read one named column out of any Ken French daily CSV as raw percent returns.

    Every daily file shares one shape: explanatory preamble, a header row that
    starts with a comma and names the columns, eight-digit date rows, then
    (in some files) an annual block. The header is matched on the requested
    column so the same reader serves the 3-factor files and the momentum files,
    whose column is ``Mom`` for the US series and ``WML`` for the regional ones.
    Reading stops at the first non-date token; the caller converts percent to
    decimal.
    """
    lines = raw_csv.splitlines()
    header_idx: int | None = None
    value_idx: int | None = None
    for i, line in enumerate(lines):
        if not line.startswith(","):
            continue
        parts = [part.strip() for part in line.split(",")]
        if column in parts:
            header_idx = i
            value_idx = parts.index(column)
            break
    if header_idx is None or value_idx is None:
        return pd.DataFrame(columns=["date", "ret"])
    rows: list[dict[str, str]] = []
    for line in lines[header_idx + 1 :]:
        parts = [part.strip() for part in line.split(",")]
        if len(parts) <= value_idx:
            continue
        date_tok = parts[0]
        if not (len(date_tok) == 8 and date_tok.isdigit()):
            break
        rows.append({"date": date_tok, "ret": parts[value_idx]})
    return pd.DataFrame(rows, columns=["date", "ret"])


def _usa_series(rows: Iterator[tuple[object, ...]]) -> tuple[list[object], list[object]] | None:
    """(dates, USA returns) of one sheet, or None when it has no DATE header with a USA column.

    The first useful row is the DATE header with a USA column. After next(), iteration
    resumes at the first data row.
    """
    header = next(
        (cells for cells in rows if cells and cells[0] == "DATE" and "USA" in cells),
        None,
    )
    if header is None:
        return None
    usa = header.index("USA")
    dates: list[object] = []
    rets: list[object] = []
    for cells in rows:  # resumes on the first data row after the header
        if not cells or cells[0] in (None, "") or usa >= len(cells):
            continue
        value = cells[usa]
        if value is None or value == "":
            continue  # AQR has no USA value for this date
        dates.append(cells[0])
        rets.append(value)
    return dates, rets


def _daily_frame(dates: list[object], rets: list[object]) -> pd.DataFrame:
    stamp = pd.to_datetime(pd.Series(dates), format="%m/%d/%Y", errors="coerce")
    if stamp.isna().all():  # date cells may be real datetimes instead of text
        stamp = pd.to_datetime(pd.Series(dates), errors="coerce")
    out = pd.DataFrame({"date": stamp, "ret": rets})
    return out.dropna(subset=["date"]).reset_index(drop=True)


def _parse_aqr_daily(payload: bytes) -> pd.DataFrame:
    """Parse one AQR daily factor workbook into ``date``/``ret``.

    The useful sheet has disclaimer rows before a ``DATE`` header with country columns. We take the
    ``USA`` column to match the US Ken French factors. AQR values are already decimal returns.
    """
    import io

    try:
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001, fail open on a missing engine or non-XLSX payload
        log.warning("aqr: could not open workbook: %s", exc)
        return pd.DataFrame(columns=["date", "ret"])
    try:
        for sheet_name in workbook.sheetnames:
            series = _usa_series(workbook[sheet_name].iter_rows(values_only=True))
            if series is None:
                continue  # definition or disclaimer sheet
            dates, rets = series
            if not dates:
                return pd.DataFrame(columns=["date", "ret"])
            return _daily_frame(dates, rets)
        log.warning("aqr: no DATE/USA factor sheet (sheets=%s)", workbook.sheetnames)
        return pd.DataFrame(columns=["date", "ret"])
    except Exception as exc:  # noqa: BLE001, fail open on any parse fault
        log.warning("aqr: parse failed: %s", exc)
        return pd.DataFrame(columns=["date", "ret"])
    finally:
        workbook.close()


def live_kenfrench_archive(archive: str) -> KenFrenchZip:
    """Build the live ZIP transport for one named Ken French daily archive."""
    url = _ARCHIVE_URLS[archive]
    return lambda: fetch_url(Provider.KEN_FRENCH, url)


@dataclass(frozen=True)
class KenFrenchFactorReturnsClient:
    """Ken French factor-return client, one transport per daily archive.

    ``transports`` is keyed by archive (``US_FF3``, ``US_MOMENTUM``, …); a factor
    id whose archive is absent from the mapping fetches empty rather than
    raising, so a partially configured client degrades one factor at a time.
    """

    transports: Mapping[str, KenFrenchZip]

    @property
    def provider(self) -> Provider:
        return Provider.KEN_FRENCH

    def fetch(self, req: FetchRequest) -> pd.DataFrame:
        if req.dataset != _DATASET:
            return pd.DataFrame()
        factor_id = req.provider_symbol
        try:
            long = self._long_for(factor_id)
        except Exception as exc:  # noqa: BLE001, the fetch fails open
            log.warning("kenfrench factor_returns: fetch %s failed: %s", factor_id, exc)
            return pd.DataFrame(columns=["factor_id", "date", "ret"])
        return _window(long, req)

    def _long_for(self, factor_id: str) -> pd.DataFrame:
        series = _KENFRENCH_SERIES.get(factor_id)
        if series is None:
            return pd.DataFrame(columns=["factor_id", "date", "ret"])
        transport = self.transports.get(series.archive)
        if transport is None:
            log.warning(
                "kenfrench factor_returns: no transport for archive %s (factor %s)",
                series.archive,
                factor_id,
            )
            return pd.DataFrame(columns=["factor_id", "date", "ret"])
        block = _parse_column_daily(_unzip_first_member(transport()), series.column)
        return pd.DataFrame(
            {
                "factor_id": factor_id,
                "date": block["date"],
                "ret": _percent_to_decimal(block["ret"]),
            }
        )

    def datasets(self) -> tuple[str, ...]:
        return (_DATASET,)


@dataclass(frozen=True)
class AqrFactorReturnsClient:
    """AQR factor-return client for QMJ and BAB."""

    transport: AqrFactorFile

    @property
    def provider(self) -> Provider:
        return Provider.AQR

    def fetch(self, req: FetchRequest) -> pd.DataFrame:
        if req.dataset != _DATASET:
            return pd.DataFrame()
        factor_id = req.provider_symbol
        if factor_id not in ("QMJ", "BAB"):
            return pd.DataFrame(columns=["factor_id", "date", "ret"])
        try:
            parsed = _parse_aqr_daily(self.transport(factor_id))
        except Exception as exc:  # noqa: BLE001, fail open on transport or format fault
            log.warning("aqr factor_returns: fetch %s failed: %s", factor_id, exc)
            return pd.DataFrame(columns=["factor_id", "date", "ret"])
        if parsed.empty:
            return pd.DataFrame(columns=["factor_id", "date", "ret"])
        long = pd.DataFrame({"factor_id": factor_id, "date": parsed["date"], "ret": parsed["ret"]})
        return _window(long, req)

    def datasets(self) -> tuple[str, ...]:
        return (_DATASET,)


def live_aqr() -> AqrFactorFile:
    """Build the live AQR workbook transport. The browser user-agent keeps AQR's CDN from
    rejecting the request."""

    def _fetch(factor_id: str) -> bytes:
        url = _AQR_URLS.get(factor_id)
        if url is None:
            raise ValueError(f"no AQR Data Library URL for factor {factor_id!r}")
        return fetch_url(Provider.AQR, url, {"User-Agent": _AQR_USER_AGENT})

    return _fetch
